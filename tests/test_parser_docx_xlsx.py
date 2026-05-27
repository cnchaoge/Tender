from docx import Document
from openpyxl import Workbook
from server.core.parser.document import parse_docx, parse_xlsx


def test_parse_docx_with_table(tmp_path):
    p = tmp_path / 'test.docx'
    doc = Document()
    doc.add_paragraph('Intro paragraph')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'A'
    table.cell(0, 1).text = 'B'
    table.cell(1, 0).text = 'C'
    table.cell(1, 1).text = 'D'
    doc.save(str(p))

    content, paras = parse_docx(str(p))
    assert 'Intro paragraph' in content
    # Table rows should appear in content prefixed with [表格]
    assert any('[表格]' in t for t in content.split('\n'))


def test_parse_xlsx_basic(tmp_path):
    p = tmp_path / 'test.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.append(['h1', 'h2'])
    ws.append(['v1', 'v2'])
    wb.save(str(p))

    content, sheets = parse_xlsx(str(p))
    assert 'Sheet1' in content
    assert any('v1' in s['text'] for s in sheets)
